"""
LongoUZ — savollarni ommaviy import qilish (JSON / CSV / TXT / PDF).

Ikki joyda ishlatiladi:
  1) admin panel -> "Import" sahifasi (fayl yuklash, ko'rib chiqish, tasdiqlash);
  2) buyruq satri skripti `import_questions.py`.

Asosiy g'oya: faylni savollar ro'yxatiga aylantirish (parse_*), keyin har bir
savol uchun fan va mavzuni AVTOMATIK aniqlash (auto_assign) — mavzu matniga
o'xshashlik bo'yicha, savoli KAM mavzularga ustuvorlik berib.
"""

from __future__ import annotations

import csv
import io
import json
import math
import re
from collections import Counter

LETTERS = "ABCD"
MAX_ITEMS = 2500
MAX_FILE_BYTES = 25 * 1024 * 1024
MAX_ZIP_ENTRIES = 300
MAX_ZIP_UNCOMPRESSED = 120 * 1024 * 1024
SUPPORTED_EXT = (".pdf", ".docx", ".xlsx", ".json", ".csv", ".tsv", ".txt", ".md", ".html", ".htm", ".zip", ".png", ".jpg", ".jpeg", ".webp", ".gif")

# ---------------------------------------------------------------------------
# Umumiy tekshiruv
# ---------------------------------------------------------------------------


def _clean(text: str) -> str:
    return re.sub(r"[ \t]+", " ", (text or "").replace(" ", " ")).strip()


def normalize_item(raw: dict) -> dict | None:
    """Turli kalit nomlarini yagona shaklga keltiradi. Yaroqsiz bo'lsa None."""
    question = _clean(str(raw.get("question") or raw.get("savol") or raw.get("q") or ""))
    options = raw.get("options")
    if options is None:
        for keys in (("A", "B", "C", "D"), ("a", "b", "c", "d")):
            if raw.get(keys[0]) is not None:
                options = [raw.get(k) for k in keys]
                break
    if options is None:
        options = [raw.get(f"option_{i}") or raw.get(f"variant_{i}") or raw.get(f"v{i}") for i in range(1, 5)]
    options = [_clean(str(o)) for o in (options or []) if o is not None and str(o).strip()]
    if not question or len(options) != 4 or len({o.lower() for o in options}) != 4:
        return None

    ci = raw.get("correct_index", raw.get("javob_index"))
    if ci is None:
        ans = raw.get("correct") or raw.get("javob") or raw.get("answer") or raw.get("to'g'ri")
        if ans is not None:
            a = str(ans).strip()
            if a.upper() in tuple(LETTERS):
                ci = LETTERS.index(a.upper())
            elif a.isdigit() and 1 <= int(a) <= 4 and "correct_index" not in raw:
                ci = int(a) - 1
            else:
                for i, o in enumerate(options):
                    if o.lower() == a.lower():
                        ci = i
                        break
    try:
        ci = int(ci) if ci is not None else None
    except (TypeError, ValueError):
        ci = None
    if ci is not None and not 0 <= ci <= 3:
        ci = None
    return {
        "question": question,
        "options": options,
        "correct_index": ci,
        "explanation": _clean(str(raw.get("explanation") or raw.get("izoh") or "")),
        "topic": _clean(str(raw.get("topic") or raw.get("mavzu") or "")),
        "subject": _clean(str(raw.get("subject") or raw.get("fan") or "")),
        "mode": _clean(str(raw.get("mode") or raw.get("rejim") or "")),
        **({"figure": raw["figure"]} if isinstance(raw.get("figure"), dict) else {}),
    }


# ---------------------------------------------------------------------------
# JSON / CSV
# ---------------------------------------------------------------------------


def parse_json_text(text: str) -> list[dict]:
    data = json.loads(text)
    if isinstance(data, dict):
        for key in ("questions", "savollar", "items", "data"):
            if isinstance(data.get(key), list):
                data = data[key]
                break
        else:
            data = [data]
    out = [normalize_item(x) for x in data if isinstance(x, dict)]
    return [x for x in out if x]


def parse_csv_text(text: str) -> list[dict]:
    sample = text[:2000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    reader = csv.DictReader(io.StringIO(text), dialect=dialect)
    rows = []
    for row in reader:
        row = {(k or "").strip().lower(): v for k, v in row.items() if k}
        rows.append(normalize_item(row))
    return [x for x in rows if x]


# ---------------------------------------------------------------------------
# Oddiy matn / PDF (1. Savol? A) ... B) ... C) ... D) ...  + javoblar kaliti)
# ---------------------------------------------------------------------------

_Q_START = re.compile(r"(?m)^\s*(\d{1,3})\s*[\.\)]\s+(?=\S)")
_KEY_HEADER = re.compile(r"(?im)^\s*(javoblar(?:\s+kaliti)?|kalitlar|kalit|answers?(?:\s+key)?|key|ответы|ключи?)\s*(?::|\-|$)")
_KEY_PAIR = re.compile(r"(\d{1,3})\s*[\-–—:\.\)]?\s*([A-Da-dА-Гa-г])\b")
_CYR2LAT = {"А": "A", "В": "B", "С": "C", "Д": "D", "а": "A", "в": "B", "с": "C", "д": "D"}
_INLINE_ANS = re.compile(
    r"(?im)^\s*(?:to'g'ri\s+javob|to‘g‘ri\s+javob|javob|answer|correct|ответ)\s*[:\-]?\s*\(?([A-Da-d])\)?\.?\s*$"
)
_INLINE_EXP = re.compile(r"(?is)\b(?:izoh|tushuntirish|explanation|пояснение)\s*[:\-]\s*(.+)$")


def _split_options(block: str):
    """block ichidan A) B) C) D) variantlarini ketma-ket topadi. (stem, [opts]) yoki None."""
    pat = re.compile(r"(?:(?<=\s)|^)\(?([A-Da-dАВСДавсд])[\)\.]\s+")
    marks = [(m.start(), m.end(), _CYR2LAT.get(m.group(1), m.group(1)).upper()) for m in pat.finditer(block)]
    # birinchi A dan boshlab A,B,C,D ketma-ketligini qidiramiz
    for i, (_, _, ch) in enumerate(marks):
        if ch != "A":
            continue
        seq = [marks[i]]
        want = "B"
        for mk in marks[i + 1:]:
            if mk[2] == want:
                seq.append(mk)
                if want == "D":
                    break
                want = chr(ord(want) + 1)
        if len(seq) == 4:
            stem = block[:seq[0][0]]
            opts = []
            for j, mk in enumerate(seq):
                end = seq[j + 1][0] if j < 3 else len(block)
                opts.append(block[mk[1]:end])
            return stem, opts
    return None


def parse_plain_text(text: str) -> list[dict]:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    key_map: dict[int, int] = {}

    # javoblar kaliti bo'limi (oxirida)
    mk = list(_KEY_HEADER.finditer(text))
    if mk:
        head = mk[-1]
        tail = text[head.end():]
        for n, letter in _KEY_PAIR.findall(tail):
            letter = _CYR2LAT.get(letter, letter).upper()
            if letter in LETTERS:
                key_map.setdefault(int(n), LETTERS.index(letter))
        if key_map:
            text = text[:head.start()]

    starts = list(_Q_START.finditer(text))
    items: list[dict] = []
    ordered_numbers: list[int] = []
    for idx, m in enumerate(starts):
        block = text[m.end(): starts[idx + 1].start() if idx + 1 < len(starts) else len(text)]
        block = block.strip()
        # satr uzilishlarini bo'sh joyga aylantiramiz (PDF satr bo'laklari)
        # lekin variantlar yangi qatordan boshlanishi mumkin — shuning uchun oldin belgilaymiz
        inline_ans = None
        am = _INLINE_ANS.search(block)
        if am:
            inline_ans = LETTERS.index(am.group(1).upper())
            block = block[:am.start()] + block[am.end():]
        explanation = ""
        em = _INLINE_EXP.search(block)
        if em:
            explanation = _clean(em.group(1).replace("\n", " "))
            block = block[:em.start()]
        flat = re.sub(r"\s*\n\s*", " ", block)
        parsed = _split_options(flat)
        if not parsed:
            continue
        stem, opts = parsed
        number = int(m.group(1))
        item = normalize_item({
            "question": stem, "options": opts, "correct_index": inline_ans, "explanation": explanation,
        })
        if item:
            item["_number"] = number
            items.append(item)
            ordered_numbers.append(number)

    for it in items:
        if it["correct_index"] is None and it["_number"] in key_map:
            it["correct_index"] = key_map[it["_number"]]
    # raqamlar mos kelmasa, lekin kalit soni savollar soniga teng bo'lsa — tartib bo'yicha
    if key_map and any(it["correct_index"] is None for it in items) and len(key_map) == len(items):
        for it, k in zip(items, [key_map[n] for n in sorted(key_map)]):
            if it["correct_index"] is None:
                it["correct_index"] = k
    for pos, it in enumerate(items, 1):
        it["no"] = it.pop("_number", None) or pos
    return items


def parse_pdf_bytes(data: bytes) -> tuple[list[dict], str]:
    """PDF -> (savollar, ogohlantirish). Skaner qilingan (rasm) PDF bo'lsa matn chiqmaydi."""
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = []
    for page in reader.pages[:300]:
        try:
            pages.append(page.extract_text() or "")
        except Exception:  # noqa: BLE001 — bitta buzuq sahifa butun faylni to'xtatmasin
            pages.append("")
    text = "\n".join(pages)
    if len(text.strip()) < 40:
        return [], "PDF ichidan matn topilmadi (skaner qilingan rasm bo'lishi mumkin). Matnli PDF yoki JSON/CSV yuklang."
    items = parse_plain_text(text)
    if items:
        try:
            import question_figures
            found = question_figures.pdf_question_figures(data)
            if any(f for _, f in found):
                question_figures.attach_pdf_figures(items, found)
        except Exception:  # noqa: BLE001 — chizma ajratish xatosi savollarni to'xtatmasin
            pass
    if not items:
        return [], "PDF'dan savollar topilmadi. Format: '1. Savol? A) ... B) ... C) ... D) ...' va oxirida 'Javoblar: 1-A 2-C ...'."
    return items, ""


def parse_docx_bytes(data: bytes) -> tuple[list[dict], str]:
    """Word (.docx) -> savollar. Paragraflar va jadval kataklari tartib bilan o'qiladi."""
    from docx import Document

    doc = Document(io.BytesIO(data))
    lines: list[str] = []
    body = doc.element.body
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, doc)
            if child.xpath(".//a:blip"):
                import question_figures
                chunks = []
                for run in para.runs:
                    chunks.append(run.text + (" " + question_figures.docx_marker_text(run._r, doc.part) + " "))
                lines.append("".join(chunks))
            else:
                lines.append(para.text)
        elif tag == "tbl":
            for row in Table(child, doc).rows:
                cells = []
                for c in row.cells:
                    t = c.text.strip()
                    if t and (not cells or cells[-1] != t):
                        cells.append(t)
                lines.append("  ".join(cells))
    text = "\n".join(lines)
    if len(text.strip()) < 20:
        return [], "Word fayl ichida matn topilmadi."
    items = parse_plain_text(text)
    if "[[IMG:" in text:
        import question_figures
        question_figures.resolve_markers(items)
    if not items:
        return [], "Word fayldan savollar topilmadi. Format: '1. Savol? A) ... B) ... C) ... D) ...' va 'Javoblar: 1-A 2-C ...'."
    return items, ""


def parse_xlsx_bytes(data: bytes) -> tuple[list[dict], str]:
    """Excel (.xlsx): birinchi qator sarlavha (question, A, B, C, D, correct, ...)."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    items: list[dict] = []
    for ws in wb.worksheets[:5]:
        buf = io.StringIO()
        w = csv.writer(buf)
        for row in ws.iter_rows(values_only=True):
            w.writerow(["" if v is None else str(v) for v in row])
        items.extend(parse_csv_text(buf.getvalue()))
    if not items:
        return [], "Excel fayldan savollar topilmadi — sarlavha: question, A, B, C, D, correct, explanation, topic."
    return items, ""


def parse_html_text(raw: str) -> list[dict]:
    text = re.sub(r"(?is)<(script|style).*?</\1>", " ", raw)
    text = re.sub(r"(?i)<br\s*/?>|</p>|</div>|</li>|</tr>|</h\d>", "\n", text)
    text = re.sub(r"<[^>]+>", " ", text)
    import html as _html
    return parse_plain_text(_html.unescape(text))


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1251", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _parse_single(name: str, data: bytes) -> tuple[list[dict], str]:
    low = (name or "").lower()
    if len(data) > MAX_FILE_BYTES:
        return [], "Fayl juda katta (25 MB dan oshmasin)."
    try:
        if low.endswith(".pdf"):
            return parse_pdf_bytes(data)
        if low.endswith(".docx"):
            return parse_docx_bytes(data)
        if low.endswith(".doc"):
            return [], "Eski .doc format qo'llab-quvvatlanmaydi — Word'da .docx qilib saqlang."
        if low.endswith(".xlsx"):
            return parse_xlsx_bytes(data)
        text = _decode(data)
        if low.endswith(".json"):
            items, warn = parse_json_text(text), ""
        elif low.endswith((".csv", ".tsv")):
            items, warn = parse_csv_text(text), ""
        elif low.endswith((".html", ".htm")):
            items, warn = parse_html_text(text), ""
        elif low.endswith((".txt", ".md")) or "." not in low:
            items, warn = parse_plain_text(text), ""
        else:
            return [], "Bu fayl turi qo'llab-quvvatlanmaydi."
        if not items and not warn:
            warn = "Fayldan savollar topilmadi — formatni tekshiring."
        return items, warn
    except json.JSONDecodeError as exc:
        return [], f"JSON xato: {exc}"
    except Exception as exc:  # noqa: BLE001 — bitta buzuq fayl butun yuklashni to'xtatmasin
        return [], f"Faylni o'qib bo'lmadi: {exc}"


def expand_zip(data: bytes, prefix: str, depth: int = 0) -> list[tuple[str, bytes, str]]:
    """ZIP -> [(ichki yo'l, bayt, ogohlantirish)]. Zip-bomba va yo'l hujumlaridan himoyalangan."""
    import zipfile

    out: list[tuple[str, bytes, str]] = []
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile:
        return [(prefix, b"", "ZIP fayl buzilgan.")]
    infos = [i for i in zf.infolist() if not i.is_dir() and "__MACOSX" not in i.filename and not i.filename.split("/")[-1].startswith(".")]
    if len(infos) > MAX_ZIP_ENTRIES:
        out.append((prefix, b"", f"ZIP ichida juda ko'p fayl (>{MAX_ZIP_ENTRIES}); faqat dastlabkilari olindi."))
        infos = infos[:MAX_ZIP_ENTRIES]
    total = 0
    for info in infos:
        total += info.file_size
        if total > MAX_ZIP_UNCOMPRESSED or info.file_size > MAX_FILE_BYTES:
            out.append((f"{prefix}/{info.filename}", b"", "Hajmi juda katta — o'tkazib yuborildi."))
            continue
        inner = f"{prefix}/{info.filename}"
        try:
            blob = zf.read(info)
        except Exception as exc:  # noqa: BLE001
            out.append((inner, b"", f"O'qib bo'lmadi: {exc}"))
            continue
        if info.filename.lower().endswith(".zip"):
            if depth < 1:
                out.extend(expand_zip(blob, inner, depth + 1))
            continue
        if info.filename.lower().endswith(SUPPORTED_EXT) or info.filename.lower().endswith(".doc"):
            out.append((inner, blob, ""))
    return out


def parse_files(files: list[tuple[str, bytes]]) -> tuple[list[dict], list[str]]:
    """Ko'p fayl (ZIP ham) -> (savollar, ogohlantirishlar). Har savolda 'src' (fayl nomi) va 'no' (fayldagi tartib raqami)."""
    all_items: list[dict] = []
    warns: list[str] = []
    queue: list[tuple[str, bytes, str]] = []
    images: list[tuple[str, bytes]] = []
    for name, data in files:
        base = (name or "fayl").replace("\\", "/").split("/")[-1][:120]
        if base.lower().endswith(".zip"):
            queue.extend(expand_zip(data, base))
        elif base.lower().endswith((".png", ".jpg", ".jpeg", ".webp", ".gif")):
            images.append((base, data))
        else:
            queue.append((base, data, ""))
    images.extend((n, d) for n, d, _w in queue if n.lower().endswith((".png", ".jpg", ".jpeg", ".webp", ".gif")) and d)
    queue = [q for q in queue if not q[0].lower().endswith((".png", ".jpg", ".jpeg", ".webp", ".gif"))]
    for name, data, pre_warn in queue:
        if pre_warn:
            warns.append(f"{name}: {pre_warn}")
            if not data:
                continue
        items, warn = _parse_single(name, data)
        if warn:
            warns.append(f"{name}: {warn}")
        for pos, it in enumerate(items, 1):
            it["src"] = name[-150:]
            it["no"] = it.get("no") or pos
        all_items.extend(items)
        if len(all_items) >= MAX_ITEMS:
            warns.append(f"Faqat dastlabki {MAX_ITEMS} ta savol olindi.")
            all_items = all_items[:MAX_ITEMS]
            break
    if images and all_items:
        import question_figures
        attached = question_figures.attach_zip_images(all_items, question_figures.zip_image_index(images))
        if attached:
            warns.append(f"{attached} ta savolga rasm/chizma biriktirildi.")
    return all_items, warns


def norm_key(text: str) -> str:
    """Savol matnini takrorni aniqlash uchun normallashtiradi (registr, bo'shliq, tinish belgilari)."""
    t = re.sub(r"[\W_]+", " ", (text or "").lower(), flags=re.UNICODE)
    return re.sub(r"\s+", " ", t).strip()


def parse_upload(filename: str, data: bytes) -> tuple[list[dict], str]:
    """Bitta fayl (yoki ZIP) -> (savollar, ogohlantirish). Eski chaqiruvlar bilan mos."""
    items, warns = parse_files([(filename, data)])
    return items, " ".join(warns)


# ---------------------------------------------------------------------------
# Fan va mavzuni avtomatik aniqlash (kam savolli mavzularga ustuvorlik)
# ---------------------------------------------------------------------------

_TOKEN = re.compile(r"[a-zа-яёʻ'’`0-9]{3,}", re.IGNORECASE)


def _tokens(text: str) -> list[str]:
    return [t.lower().replace("’", "'").replace("`", "'") for t in _TOKEN.findall(text or "")]


class TopicMatcher:
    """TF-IDF + kosinus o'xshashlik. samples: {(subject, topic): [savol matnlari]}"""

    def __init__(self, samples: dict[tuple[str, str], list[str]], counts: dict[tuple[str, str], int], target: int = 20):
        self.keys = list(samples)
        self.counts = counts
        self.target = target
        docs = {k: Counter(t for q in v for t in _tokens(q)) for k, v in samples.items()}
        df: Counter = Counter()
        for c in docs.values():
            df.update(c.keys())
        n = max(1, len(docs))
        self.idf = {t: math.log((1 + n) / (1 + d)) + 1 for t, d in df.items()}
        self.vecs: dict[tuple[str, str], dict[str, float]] = {}
        for k, c in docs.items():
            vec = {t: (1 + math.log(f)) * self.idf[t] for t, f in c.items()}
            norm = math.sqrt(sum(v * v for v in vec.values())) or 1.0
            self.vecs[k] = {t: v / norm for t, v in vec.items()}

    def best(self, text: str, allowed_subjects: set[str] | None = None) -> tuple[tuple[str, str] | None, float]:
        c = Counter(_tokens(text))
        q = {t: (1 + math.log(f)) * self.idf.get(t, 0.0) for t, f in c.items() if t in self.idf}
        qn = math.sqrt(sum(v * v for v in q.values())) or 1.0
        best_key, best_score, best_sim = None, -1.0, 0.0
        for k, vec in self.vecs.items():
            if allowed_subjects and k[0] not in allowed_subjects:
                continue
            sim = sum(w * vec.get(t, 0.0) for t, w in q.items()) / qn
            deficit = max(0.0, (self.target - self.counts.get(k, 0)) / self.target)
            score = sim * (1 + 1.5 * deficit)
            if score > best_score:
                best_key, best_score, best_sim = k, score, sim
        return best_key, best_sim

    def thinnest(self, allowed_subjects: set[str] | None = None):
        cands = [k for k in self.keys if not allowed_subjects or k[0] in allowed_subjects]
        return min(cands, key=lambda k: self.counts.get(k, 0)) if cands else None


def auto_assign(items: list[dict], matcher: TopicMatcher, subject: str | None, force_topic: str | None = None) -> None:
    """items ichiga 'subject' va 'topic' yozadi. subject None/'auto' bo'lsa fan ham aniqlanadi."""
    allowed = {subject} if subject and subject != "auto" else None
    for it in items:
        if force_topic and subject and subject != "auto":
            it["subject"], it["topic"], it["auto"] = subject, force_topic, False
            continue
        if it.get("topic") and it.get("subject") and not allowed:
            it["auto"] = False
            continue
        text = it["question"] + " " + " ".join(it["options"])
        key, sim = matcher.best(text, allowed)
        if key is None or sim < 0.03:
            key = matcher.thinnest(allowed)
        if key:
            it["subject"], it["topic"] = key
            matcher.counts[key] = matcher.counts.get(key, 0) + 1  # keyingi savollar muvozanatli tarqalsin
        it["auto"] = True


async def build_matcher(mode: str, subjects: list[str] | None = None) -> TopicMatcher:
    """Bazadagi mavjud savollardan mavzu moslashtiruvchini quradi."""
    import database as db

    codes = subjects or [s["code"] for s in db.SUBJECTS]
    samples: dict[tuple[str, str], list[str]] = {}
    counts: dict[tuple[str, str], int] = {}
    for code in codes:
        by_topic = await db.sample_questions_for_matching(code, mode)
        totals = await db.count_questions_by_topic(code, mode)
        for topic, lst in by_topic.items():
            samples[(code, topic)] = lst
            counts[(code, topic)] = totals.get(topic, len(lst))
    return TopicMatcher(samples, counts)


async def thin_topics(mode: str, threshold: int = 15) -> list[dict]:
    """Savoli `threshold` dan kam mavzular (admin import sahifasida ko'rsatiladi)."""
    import database as db

    out = []
    for s in db.SUBJECTS:
        totals = await db.count_questions_by_topic(s["code"], mode)
        for topic, c in totals.items():
            if c < threshold:
                out.append({"subject": s["code"], "subject_name": s["name"], "topic": topic, "count": c})
    out.sort(key=lambda x: x["count"])
    return out
